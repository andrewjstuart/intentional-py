import json
import zipfile
from pathlib import Path

import openpyxl
import pytest

from intentional_py import build_intents, exceptions, extract
from intentional_py import validate as validating
from intentional_py.models import LanguageSettings, ProjectLayout
from intentional_py.tests.conftest import FakeReporter, make_nl_phrases


def test_nl_config_aborts_when_duplicates_declined(tmp_path: Path) -> None:
    config = make_nl_phrases(tmp_path)
    reporter = FakeReporter(answer=False)

    with pytest.raises(exceptions.IntentionalException):
        build_intents.nl_config(config, "RTL", "GetIntent", False, reporter)

    assert reporter.questions == ["Continue processing files"]
    assert reporter.question_details == [["hello"]]


def test_nl_build_returns_result(tmp_path: Path) -> None:
    config = make_nl_phrases(tmp_path)
    reporter = FakeReporter(answer=True)

    build_intents.nl_config(config, "RTL", "GetIntent", False, reporter)
    result = build_intents.intents("NL", config, tmp_path, reporter)

    assert result.intents == 2
    assert result.phrases == 3
    assert result.files == 4
    assert result.languages == ["en"]
    assert (tmp_path / "intents" / "RTL.Billing.json").exists()


def test_nl_config_requires_training_phrases(tmp_path: Path) -> None:
    with pytest.raises(exceptions.FileSystemError):
        build_intents.nl_config(
            tmp_path / "intents_nl.cfg", "RTL", "GetIntent", False, FakeReporter()
        )


def test_nl_cli_uses_config_directory(tmp_path: Path, monkeypatch) -> None:
    from typer.testing import CliRunner

    from intentional_py.intentional import app

    project = tmp_path / "project"
    config = make_nl_phrases(project)
    elsewhere = tmp_path / "elsewhere"
    elsewhere.mkdir()
    monkeypatch.chdir(elsewhere)

    result = CliRunner().invoke(
        app, ["nl", "--config", str(config), "-v", "RTL"], input="y\n"
    )

    assert result.exit_code == 0, result.stdout
    assert config.exists()
    assert (project / "intents" / "RTL.Billing.json").exists()
    assert not (elsewhere / "intents").exists()


def test_dd_build_machine_learning_column(tmp_path: Path) -> None:
    phrase_dir = tmp_path / "Training Phrases" / "en"
    phrase_dir.mkdir(parents=True)
    (phrase_dir / "on.txt").write_text("turn it on\n", encoding="utf-8")
    (phrase_dir / "off.txt").write_text("turn it off\n", encoding="utf-8")
    config = tmp_path / "intents.cfg"
    config.write_text(
        "A.Default,Ctx,en,on,,,\nA.On,Ctx,en,on,,,TRUE\nA.Off,Ctx,en,off,,,FALSE\n",
        encoding="utf-8",
    )

    result = build_intents.intents("DD", config, tmp_path, FakeReporter())

    intents_dir = tmp_path / "intents"
    assert (
        json.loads((intents_dir / "A.Default.json").read_text(encoding="utf-8"))["auto"]
        is True
    )
    assert (
        json.loads((intents_dir / "A.On.json").read_text(encoding="utf-8"))["auto"]
        is True
    )
    assert (
        json.loads((intents_dir / "A.Off.json").read_text(encoding="utf-8"))["auto"]
        is False
    )
    assert result.machine_learning_off == ["A.Off"]


def test_dd_build_with_a_custom_project_layout(tmp_path: Path) -> None:
    layout = ProjectLayout(training_phrases_dir="Phrases", intents_dir="output")
    phrase_dir = tmp_path / "Phrases" / "en"
    phrase_dir.mkdir(parents=True)
    (phrase_dir / "billing.txt").write_text("pay my bill\n", encoding="utf-8")
    config = tmp_path / "intents.cfg"
    config.write_text("MYAC.Billing,Ctx,en,billing,,,\n", encoding="utf-8")

    result = build_intents.intents(
        "DD", config, tmp_path, FakeReporter(), layout=layout
    )

    assert result.output_dir == tmp_path / "output"
    assert (tmp_path / "output" / "MYAC.Billing_usersays_en.json").exists()
    assert not (tmp_path / "intents").exists()
    # the default folder names are not used at all with a custom layout
    assert not (tmp_path / "Training Phrases").exists()


def test_nl_build_with_a_custom_project_layout(tmp_path: Path) -> None:
    layout = ProjectLayout(training_phrases_dir="Phrases", nl_subfolder="Natural")
    nl_dir = tmp_path / "Phrases" / "en" / "Natural"
    nl_dir.mkdir(parents=True)
    (nl_dir / "BILLING.txt").write_text("pay my bill\n", encoding="utf-8")
    config = tmp_path / "intents_nl.cfg"
    reporter = FakeReporter(answer=True)

    build_intents.nl_config(config, "RTL", "GetIntent", False, reporter, layout)
    result = build_intents.intents("NL", config, tmp_path, reporter, layout=layout)

    assert result.intents == 1
    assert (tmp_path / "intents" / "RTL.Billing.json").exists()


def test_extract_with_a_custom_project_layout(tmp_path: Path) -> None:
    layout = ProjectLayout(training_phrases_dir="Phrases", nl_subfolder="Natural")
    workbook = openpyxl.Workbook()
    workbook.active.title = "HELLO"
    workbook.active["A1"] = "hi there"
    xl_file = tmp_path / "phrases.xlsx"
    workbook.save(xl_file)

    result = extract.excel_data(xl_file, "NL", "en", tmp_path, FakeReporter(), layout)

    assert result.output_dir == tmp_path / "Phrases" / "en" / "Natural"
    assert not (tmp_path / "Training Phrases").exists()


def test_dd_build_with_a_non_default_language_set(tmp_path: Path) -> None:
    languages = LanguageSettings(
        languages={"en": "English", "de": "German"}, default_language="de"
    )
    phrase_dir = tmp_path / "Training Phrases" / "de"
    phrase_dir.mkdir(parents=True)
    (phrase_dir / "billing.txt").write_text(
        "bezahle meine Rechnung\n", encoding="utf-8"
    )
    config = tmp_path / "intents.cfg"
    config.write_text("A.Billing,Ctx,de,billing,,,\n", encoding="utf-8")

    result = build_intents.intents(
        "DD", config, tmp_path, FakeReporter(), languages=languages
    )

    assert result.intents == 1
    assert (tmp_path / "intents" / "A.Billing.json").exists()
    assert (tmp_path / "intents" / "A.Billing_usersays_de.json").exists()


def test_dd_build_synthesizes_the_configured_default_language(tmp_path: Path) -> None:
    # a row in a non-default language still needs a 'de' row with a custom default,
    # the same way one needed an 'en' row with the original hardcoded default
    languages = LanguageSettings(
        languages={"en": "English", "de": "German"}, default_language="de"
    )
    phrase_dir = tmp_path / "Training Phrases" / "en"
    phrase_dir.mkdir(parents=True)
    (phrase_dir / "billing.txt").write_text("pay my bill\n", encoding="utf-8")
    config = tmp_path / "intents.cfg"
    config.write_text("A.Billing,Ctx,en,billing,,,\n", encoding="utf-8")

    build_intents.intents("DD", config, tmp_path, FakeReporter(), languages=languages)

    assert (tmp_path / "intents" / "A.Billing_usersays_en.json").exists()


def test_validate_with_a_non_default_language_set(tmp_path: Path) -> None:
    languages = LanguageSettings(
        languages={"en": "English", "de": "German"}, default_language="de"
    )
    (tmp_path / "Training Phrases" / "de").mkdir(parents=True)
    config = tmp_path / "intents.cfg"
    config.write_text("A.Billing,Ctx,de,billing,,,\n", encoding="utf-8")

    result = validating.validate(config, tmp_path, FakeReporter(), languages=languages)

    labels = [(c.label, c.ok) for c in result.directories]
    assert ("German path", True) in labels
    # only languages actually used in the config are checked - a supported but
    # unused language doesn't need a folder yet
    assert not any(label == "English path" for label, _ok in labels)


def test_dd_cli_uses_config_directory(tmp_path: Path, monkeypatch) -> None:
    from typer.testing import CliRunner

    from intentional_py.intentional import app

    project = tmp_path / "project"
    phrase_dir = project / "Training Phrases" / "en"
    phrase_dir.mkdir(parents=True)
    (phrase_dir / "billing.txt").write_text("pay my bill\n", encoding="utf-8")
    config = project / "intents.cfg"
    config.write_text("MYAC.Billing,MYAC-Billing,en,billing,,,TRUE\n", encoding="utf-8")
    elsewhere = tmp_path / "elsewhere"
    elsewhere.mkdir()
    monkeypatch.chdir(elsewhere)

    result = CliRunner().invoke(app, ["--config", str(config)])

    assert result.exit_code == 0, result.stdout
    assert (project / "intents" / "MYAC.Billing_usersays_en.json").exists()
    assert not (elsewhere / "intents").exists()


def test_extract_reports_empty_sheets(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    workbook.active.title = "HELLO"
    workbook.active["A1"] = "hi there"
    workbook.create_sheet("EMPTY_ONE")
    xl_file = tmp_path / "phrases.xlsx"
    workbook.save(xl_file)

    result = extract.excel_data(xl_file, "NL", "en", tmp_path, FakeReporter())

    assert result.files == 2
    assert result.phrases == 1
    assert result.empty_sheets == ["EMPTY_ONE"]
    assert result.output_dir == tmp_path / "Training Phrases" / "en" / "NL"


def make_workbook(path: Path, sheet: str, phrase: str) -> Path:
    workbook = openpyxl.Workbook()
    workbook.active.title = sheet
    workbook.active["A1"] = phrase
    workbook.save(path)
    return path


def test_dd_extract_keeps_nl_phrases(tmp_path: Path) -> None:
    lang_dir = tmp_path / "Training Phrases" / "en"
    (lang_dir / "NL").mkdir(parents=True)
    (lang_dir / "NL" / "HELLO.txt").write_text("hello\n", encoding="utf-8")
    (lang_dir / "old.txt").write_text("old phrase\n", encoding="utf-8")
    xl_file = make_workbook(tmp_path / "dd.xlsx", "home", "at home")

    result = extract.excel_data(xl_file, "DD", "en", tmp_path, FakeReporter())

    assert (lang_dir / "NL" / "HELLO.txt").read_text(encoding="utf-8") == "hello\n"
    assert (lang_dir / "home.txt").read_text(encoding="utf-8") == "at home\n"
    assert not (lang_dir / "old.txt").exists()
    assert (
        result.backup is not None
        and result.backup.parent == tmp_path / "Training Phrases"
    )
    with zipfile.ZipFile(result.backup) as backup:
        assert backup.namelist() == ["old.txt"]


def test_corrupt_workbook_leaves_existing_phrases(tmp_path: Path) -> None:
    nl_dir = tmp_path / "Training Phrases" / "en" / "NL"
    nl_dir.mkdir(parents=True)
    (nl_dir / "HELLO.txt").write_text("hello\n", encoding="utf-8")
    xl_file = tmp_path / "broken.xlsx"
    xl_file.write_bytes(b"not an excel file")

    with pytest.raises(
        exceptions.ExtractionError, match="could not be read as an Excel file"
    ):
        extract.excel_data(xl_file, "NL", "en", tmp_path, FakeReporter())

    assert (nl_dir / "HELLO.txt").exists()
    assert not list(nl_dir.parent.glob("*.zip"))


def test_non_utf8_phrase_file_names_the_file(tmp_path: Path) -> None:
    config = make_nl_phrases(tmp_path)
    bad_file = tmp_path / "Training Phrases" / "en" / "NL" / "SPANISH.txt"
    bad_file.write_bytes("año\n".encode("cp1252"))

    with pytest.raises(
        exceptions.FileSystemError, match="SPANISH.txt is not saved as UTF-8"
    ):
        build_intents.nl_config(config, "RTL", "GetIntent", False, FakeReporter())
