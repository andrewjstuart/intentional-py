import io
import json
import zipfile
from pathlib import Path

import openpyxl
import pytest

from intentional_py import build_intents, exceptions
from intentional_py.tests.conftest import FakeReporter, dd_project, make_nl_phrases
from intentional_py.web.actions import (
    build_dd_project,
    build_nl_project,
    compare_project,
    design_project,
    download_project,
    extract_project,
    new_project,
    open_project,
    project_files,
    validate_project,
)


def _zip_dir(directory: Path) -> bytes:
    """Zip a directory's contents at the archive root, like a user's project zip."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for path in directory.rglob("*"):
            if path.is_file():
                archive.write(path, arcname=path.relative_to(directory))
    return buffer.getvalue()


def _archive() -> zipfile.ZipFile:
    return zipfile.ZipFile(io.BytesIO(download_project()))


def _design_workbook_bytes() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Design"
    headers = [
        "Intent",
        "Context",
        "Language",
        "Action",
        "Entities",
        "DTMF",
        "Machine Learning",
    ]
    for column, header in enumerate(headers, start=1):
        sheet.cell(1, column, header)
    sheet.cell(2, 1, "A.Pay")
    sheet.cell(2, 2, "Billing")
    sheet.cell(2, 3, "en")
    sheet.cell(2, 4, "pay")
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_a_task_without_a_project_is_rejected() -> None:
    import intentional_py.web.actions as actions_module

    actions_module._workspace = None
    with pytest.raises(exceptions.ConfigurationError):
        validate_project()


def test_open_project_rejects_a_zip_slip_entry() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("../../evil.cfg", "A.Pay,Ctx,en,pay,,,")
    new_project()

    with pytest.raises(exceptions.ConfigurationError):
        open_project(buffer.getvalue())


def test_validate_project_passes_for_a_valid_project(tmp_path: Path) -> None:
    project = tmp_path / "project"
    dd_project(
        project,
        ["A.Pay,Ctx,en,pay,,,"],
        {"pay": "pay my bill"},
    )
    open_project(_zip_dir(project))

    result = json.loads(validate_project())

    assert result["used_standard_configs"] is True
    # only English phrases exist in this project; other languages are expected to fail
    assert ("English path", True) in [tuple(d) for d in result["directories"]]
    # only intents.cfg exists; intents_nl.cfg is one of the two standard files checked
    assert ("Checking for intents.cfg", True) in [
        tuple(c) for c in result["config_files"]
    ]
    assert all(config["ok"] for config in result["configs"])


def test_validate_project_reports_row_errors(tmp_path: Path) -> None:
    project = tmp_path / "project"
    dd_project(
        project,
        ["A-B,Ctx,en,pay,,,"],
        {"pay": "pay my bill"},
    )
    open_project(_zip_dir(project))

    result = json.loads(validate_project())

    config = next(c for c in result["configs"] if not c["ok"])
    assert any("A-B" in detail for detail in config["details"])


def test_validate_project_uses_the_given_config_name(tmp_path: Path) -> None:
    project = tmp_path / "project"
    dd_project(
        project,
        ["A.Pay,Ctx,en,pay,,,"],
        {"pay": "pay my bill"},
        layout=None,
    )
    (project / "custom.cfg").write_text(
        (project / "intents.cfg").read_text(encoding="utf-8"), encoding="utf-8"
    )
    open_project(_zip_dir(project))

    result = json.loads(validate_project("custom.cfg"))

    assert result["used_standard_configs"] is False
    assert all(config["ok"] for config in result["configs"])


def test_build_dd_project_writes_into_the_same_project(tmp_path: Path) -> None:
    project = tmp_path / "project"
    dd_project(project, ["A.Pay,Ctx,en,pay,,,"], {"pay": "pay my bill"})
    open_project(_zip_dir(project))

    result = json.loads(build_dd_project())

    assert ["Intents", "1"] in result["summary"]
    assert "intents/A.Pay.json" in json.loads(project_files())["files"]
    assert "intents/A.Pay.json" in _archive().namelist()


def test_build_nl_project_requires_a_vertical(tmp_path: Path) -> None:
    project = tmp_path / "project"
    make_nl_phrases(project)
    open_project(_zip_dir(project))

    with pytest.raises(exceptions.ConfigurationError):
        build_nl_project()


def test_build_nl_project_builds_from_phrase_files(tmp_path: Path) -> None:
    project = tmp_path / "project"
    make_nl_phrases(project)
    open_project(_zip_dir(project))

    result = json.loads(build_nl_project(vertical="RTL"))

    names = _archive().namelist()
    assert any(name.startswith("intents/RTL.") for name in names)
    assert result["summary"]


def test_extract_project_writes_phrase_files_into_the_project(data_dir: Path) -> None:
    excel_bytes = (data_dir / "NL_English_Data.xlsm").read_bytes()
    new_project()

    result = json.loads(
        extract_project(excel_bytes, "NL_English_Data.xlsm", "NL", "en")
    )

    assert result["files"] > 0
    files = json.loads(project_files())["files"]
    assert any(f.startswith("Training Phrases/en/NL/") for f in files)


def test_design_project_creates_a_config_file(tmp_path: Path) -> None:
    new_project()

    result = json.loads(design_project(_design_workbook_bytes(), "design.xlsx"))

    assert result["rows"] == 1
    assert json.loads(project_files())["files"] == ["intents.cfg"]


def test_compare_project_reports_differences(tmp_path: Path) -> None:
    project = tmp_path / "project"
    config = dd_project(project, ["A.One,Ctx,en,one,,,"], {"one": "one"})
    build_intents.intents("DD", config, project, FakeReporter())
    export_buffer = io.BytesIO()
    with zipfile.ZipFile(export_buffer, "w") as archive:
        for path in (project / "intents").glob("*.json"):
            archive.write(path, arcname=f"intents/{path.name}")
    config.write_text("A.One,Ctx,en,one,,,\nA.Two,Ctx,en,one,,,\n", encoding="utf-8")
    open_project(_zip_dir(project))

    result = json.loads(compare_project(export_buffer.getvalue(), "export.zip"))

    assert result["added"] == ["A.Two"]
    assert result["unchanged"] == 1


def test_extract_then_build_nl_without_reuploading(data_dir: Path) -> None:
    """The exact workflow that used to need a download/re-upload in between."""
    excel_bytes = (data_dir / "NL_English_Data.xlsm").read_bytes()
    new_project()

    extract_result = json.loads(
        extract_project(excel_bytes, "NL_English_Data.xlsm", "NL", "en")
    )
    assert extract_result["files"] > 0

    build_result = json.loads(build_nl_project(vertical="RTL"))

    assert build_result["summary"]
    assert any(name.startswith("intents/RTL.") for name in _archive().namelist())
