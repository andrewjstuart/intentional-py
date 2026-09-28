"""Extract training phrases from Excel files.

This module handles reading Excel files (.xlsb, .xlsm, .xlsx) and extracting
training phrases into organized text files by language and mode (DD or NL).
Includes backup functionality for existing phrase directories.
"""

import datetime
import shutil
import zipfile
from pathlib import Path
from time import perf_counter

import openpyxl
import pyxlsb
from openpyxl.utils.exceptions import InvalidFileException

from intentional_py import constants, exceptions, utils
from intentional_py.reporting import ExtractResult, Reporter


def _backup_existing(phrase_dir: Path, mode: str) -> Path | None:
    """Zip and remove the phrases about to be replaced; returns the zip path if one was made."""
    if not phrase_dir.exists():
        return None
    # DD phrases share the language folder with the NL subfolder, which must be kept
    if mode == "DD":
        old_files = sorted(phrase_dir.glob(f"*{constants.PHRASE_FILE_EXTENSION}"))
    else:
        old_files = sorted(phrase_dir.rglob("*"))
    if not old_files:
        return None
    stamp = datetime.datetime.now(datetime.UTC).astimezone().strftime("%Y-%m-%d_%H%M%S")
    zip_path = Path(phrase_dir.parent, f"{phrase_dir.name}_{stamp}.zip")
    with utils.file_errors(zip_path):
        utils.zip_directory(phrase_dir, zip_path, old_files)
        if mode == "DD":
            for old_file in old_files:
                old_file.unlink()
        else:
            shutil.rmtree(phrase_dir)
    return zip_path


def excel_data(
    xl_file: Path, mode: str, language: str, base_dir: Path, reporter: Reporter
) -> ExtractResult:
    """Save each sheet's phrases as a text file under base_dir's Training Phrases folder."""
    t1_start = perf_counter()

    xl_file_path: Path
    file_extension: str

    (xl_file_path, xl_file, file_extension) = utils.check_for_path(xl_file)

    if file_extension not in constants.SUPPORTED_EXCEL_EXTENSIONS:
        raise exceptions.ExtractionError(
            f"Unsupported file extension: {file_extension}. Supported: {constants.SUPPORTED_EXCEL_EXTENSIONS}"
        )

    # set up path to save files
    # set phrase file path
    phrase_file_path: Path = Path(
        base_dir, constants.DEFAULT_TRAINING_PHRASES_DIR, language
    )
    phrase_file_path = (
        Path(phrase_file_path, "NL") if mode == "NL" else phrase_file_path
    )

    reporter.message("info", f"Exporting data to: [blue]{phrase_file_path}[/blue]")

    # read the whole workbook before touching existing phrases, so a bad file changes nothing
    label = f"Processing [green]{xl_file}[/green]"
    phrase_dict: dict = {}
    phrase_file: Path
    xl = Path(xl_file_path, xl_file)
    try:
        with utils.file_errors(xl):
            phrase_dict = _read_workbook(xl, file_extension, phrase_file_path, label, reporter)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, ValueError) as error:
        raise exceptions.ExtractionError(
            f"{xl} could not be read as an Excel file ({error})."
        ) from error

    result = ExtractResult(output_dir=phrase_file_path)
    result.backup = _backup_existing(phrase_file_path, mode)
    phrase_file_path.mkdir(parents=True, exist_ok=True)
    for phrase_file, phrases in phrase_dict.items():
        result.files += 1
        if not phrases:
            result.empty_sheets.append(phrase_file.stem)
        with utils.file_errors(phrase_file), phrase_file.open(mode="w", encoding="utf-8") as f:
            for line in phrases:
                result.phrases += 1
                f.write(f"{line}\n")

    result.elapsed = perf_counter() - t1_start
    return result


def _read_workbook(
    xl: Path, file_extension: str, phrase_file_path: Path, label: str, reporter: Reporter
) -> dict:
    """Return the sorted, de-duplicated phrases of each sheet, keyed by the phrase file to write."""
    phrase_dict: dict = {}
    if file_extension == ".xlsb":
        # uses pyxlsb
        with pyxlsb.open_workbook(xl) as wb:
            for sheet in reporter.track(wb.sheets, label):
                phrases: set = set()
                for row in wb.get_sheet(sheet).rows():
                    for cell in row:
                        # add the phrases to a set to remove duplicates
                        if cell.v is not None and str(cell.v).strip():
                            phrases.add(str(cell.v))

                phrase_dict[Path(phrase_file_path, f"{sheet}.txt")] = sorted(phrases)
    elif file_extension in [".xlsm", ".xlsx"]:
        # uses openpyxl
        wb = openpyxl.load_workbook(xl)
        for sheet_name in reporter.track(wb.sheetnames, label):
            phrases: set = set()
            sheet = wb[sheet_name]
            for row in sheet.iter_rows(values_only=True):
                phrase = "".join(str(cell) for cell in row if cell is not None)
                # add the phrases to a set to remove duplicates
                if phrase.strip():
                    phrases.add(phrase)
            phrase_dict[Path(phrase_file_path, f"{sheet_name}.txt")] = sorted(phrases)
    else:
        raise exceptions.ExtractionError(f"[red]Unsupported file format[/red]: {xl}")
    return phrase_dict
